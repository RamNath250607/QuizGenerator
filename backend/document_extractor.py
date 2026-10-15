import io
from typing import Dict, Any, List
import pymupdf
from docx import Document


class DocumentExtractionResult:
    def __init__(self, filename: str, full_text: str, pages: List[Dict[str, Any]]):
        self.filename = filename
        self.full_text = full_text.strip()
        self.pages = pages
        self.char_count = len(self.full_text)
        self.word_count = len(self.full_text.split())
        self.page_count = len(pages) if pages else 1

    def to_dict(self) -> Dict[str, Any]:
        return {
            "filename": self.filename,
            "char_count": self.char_count,
            "word_count": self.word_count,
            "page_count": self.page_count,
            "preview": self.full_text[:300] + "..." if len(self.full_text) > 300 else self.full_text,
            "page_chunks_count": len(self.pages),
        }

def extract_text_from_pdf(file_bytes: bytes, filename: str) -> DocumentExtractionResult:
    """Extract text page-by-page from PDF using PyMuPDF"""
    doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    pages_data = []

    full_text_list = []

    for page_idx in range(len(doc)):
        page = doc[page_idx]
        page_text = page.get_text("text").strip()
        if page_text:
            pages_data.append({
                "page_number": page_idx + 1,
                "text": page_text
            })
            full_text_list.append(f"--- Page {page_idx + 1} ---\n{page_text}")

    full_text = "\n\n".join(full_text_list)
    return DocumentExtractionResult(filename, full_text, pages_data)

def extract_text_from_docx(file_bytes: bytes, filename: str) -> DocumentExtractionResult:
    """Extract text from Word .docx file using python-docx"""
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = []
    
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            paragraphs.append(txt)

    # Also extract text from any tables inside docx
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                paragraphs.append(" | ".join(row_text))

    full_text = "\n\n".join(paragraphs)
    # Treat whole docx as 1 page unit for chunking
    pages_data = [{"page_number": 1, "text": full_text}]
    return DocumentExtractionResult(filename, full_text, pages_data)

def extract_text_from_txt(file_bytes: bytes, filename: str) -> DocumentExtractionResult:
    """Extract text from plain .txt file with UTF-8 / Latin-1 fallback"""
    try:
        full_text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        full_text = file_bytes.decode("latin-1", errors="replace")

    pages_data = [{"page_number": 1, "text": full_text}]
    return DocumentExtractionResult(filename, full_text, pages_data)

def extract_document(file_bytes: bytes, filename: str) -> DocumentExtractionResult:
    """Dispatch document extraction based on extension"""
    ext = filename.split(".")[-1].lower() if "." in filename else ""

    if ext == "pdf":
        return extract_text_from_pdf(file_bytes, filename)
    elif ext == "docx":
        return extract_text_from_docx(file_bytes, filename)
    elif ext == "txt":
        return extract_text_from_txt(file_bytes, filename)
    else:
        raise ValueError(f"Unsupported file format: '.{ext}'. Supported formats are PDF, DOCX, and TXT.")
